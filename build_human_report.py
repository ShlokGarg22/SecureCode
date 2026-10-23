#!/usr/bin/env python3
"""
Build the human-style, formal, emoji-free, comprehensive SecureCode Project Report.
Includes embedded visual architecture diagrams, deep technical explanations,
clean structured tables, and balanced team member contributions.
Team:
  - Shlok Garg (PRN: 24070122187)
  - Shashank Tai (PRN: 24070122181)
  - Sayantan Nandi (PRN: 24070122176)
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

# Palette: Traditional Academic / Engineering Report (Navy & Slate)
COLOR_NAVY = RGBColor(27, 54, 93)       # #1B365D Primary Heading
COLOR_SLATE = RGBColor(51, 65, 85)      # #334155 Secondary Heading
COLOR_BODY = RGBColor(30, 41, 59)       # #1E293B Body Text
COLOR_MUTED = RGBColor(100, 116, 139)   # #64748B Captions / Headers
COLOR_LINE = "CBD5E1"                   # Table Borders
HEX_HEADER_BG = "1B365D"
HEX_ROW_ALT = "F8FAFC"
HEX_CODE_BG = "F1F5F9"


def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    """Set cell internal padding."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def set_cell_shading(cell, color_hex):
    """Set cell background shading."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)


def set_cell_borders(cell, top="none", bottom="none", left="none", right="none", color="CBD5E1", sz="6"):
    """Set individual cell border lines."""
    tcPr = cell._tc.get_or_add_tcPr()
    borders_xml = f'<w:tcBorders {nsdecls("w")}>'
    for side, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        if val != "none":
            borders_xml += f'<w:{side} w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        else:
            borders_xml += f'<w:{side} w:val="none"/>'
    borders_xml += '</w:tcBorders>'
    tcPr.append(parse_xml(borders_xml))


def style_table_header(row, col_widths, titles):
    """Style table header with navy background and white text."""
    for idx, cell in enumerate(row.cells):
        cell.width = Inches(col_widths[idx])
        set_cell_shading(cell, HEX_HEADER_BG)
        set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(titles[idx])
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(255, 255, 255)


def add_row(table, col_widths, values, is_alt=False, is_bold_first=False):
    """Add a row of data with alternating shading."""
    row = table.add_row()
    for idx, cell in enumerate(row.cells):
        cell.width = Inches(col_widths[idx])
        bg = HEX_ROW_ALT if is_alt else "FFFFFF"
        set_cell_shading(cell, bg)
        set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
        set_cell_borders(cell, bottom="single", color=COLOR_LINE, sz="4")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(str(values[idx]))
        run.font.name = "Arial"
        run.font.size = Pt(9)
        run.font.color.rgb = COLOR_BODY
        if idx == 0 and is_bold_first:
            run.bold = True
    return row


def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(36)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(26)
    run.font.color.rgb = COLOR_NAVY
    return p


def add_subtitle(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(24)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(12)
    run.font.italic = True
    run.font.color.rgb = COLOR_SLATE
    return p


def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(14)
    run.font.color.rgb = COLOR_NAVY
    return p


def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(11.5)
    run.font.color.rgb = COLOR_SLATE
    return p


def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(10)
    run.font.color.rgb = COLOR_BODY
    return p


def add_p(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(9.5)
    run.font.color.rgb = COLOR_BODY
    return p


def add_bullet(doc, lead, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    r_lead = p.add_run(f"{lead}: ")
    r_lead.bold = True
    r_lead.font.name = "Arial"
    r_lead.font.size = Pt(9.5)
    r_lead.font.color.rgb = COLOR_BODY
    
    r_text = p.add_run(text)
    r_text.font.name = "Arial"
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = COLOR_BODY
    return p


def add_code(doc, code_str, caption=None):
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.paragraph_format.space_before = Pt(6)
        p_cap.paragraph_format.space_after = Pt(2)
        r_cap = p_cap.add_run(caption)
        r_cap.font.name = "Arial"
        r_cap.font.size = Pt(8.5)
        r_cap.font.italic = True
        r_cap.font.color.rgb = COLOR_MUTED

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_shading(cell, HEX_CODE_BG)
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
    set_cell_borders(cell, top="single", bottom="single", left="single", right="single", color=COLOR_LINE, sz="4")
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(code_str.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(8)
    run.font.color.rgb = COLOR_BODY
    
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(4)


def add_image_figure(doc, img_path, caption):
    """Embed an image with centered alignment and formal figure caption."""
    if not os.path.exists(img_path):
        print(f"[!] Warning: Image {img_path} not found.")
        return

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(img_path, width=Inches(6.2))

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(10)
    r_cap = p_cap.add_run(caption)
    r_cap.font.name = "Arial"
    r_cap.font.size = Pt(8.5)
    r_cap.font.italic = True
    r_cap.font.color.rgb = COLOR_MUTED


def build_report():
    print("[*] Generating human-styled, emoji-free SecureCode Project Report...")
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        footer = section.footer
        p_ft = footer.paragraphs[0]
        p_ft.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_ft = p_ft.add_run("SecureCode Project Report | Academic & Technical Documentation")
        r_ft.font.name = "Arial"
        r_ft.font.size = Pt(8)
        r_ft.font.color.rgb = COLOR_MUTED

    # -------------------------------------------------------------
    # COVER / HEADER BLOCK
    # -------------------------------------------------------------
    add_title(doc, "SecureCode")
    add_subtitle(doc, "Zero-Hosting Automated CI/CD Security Scanner & AI-Powered Remediation Platform")

    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.paragraph_format.space_after = Pt(18)
    r_div = p_div.add_run("-" * 65)
    r_div.font.color.rgb = COLOR_SLATE

    # Authors Table
    auth_table = doc.add_table(rows=4, cols=3)
    auth_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    w_auth = [2.1, 2.0, 2.4]
    style_table_header(auth_table.rows[0], w_auth, ["Team Member", "PRN Number", "Core Area of Ownership"])
    add_row(auth_table, w_auth, ["Shlok Garg", "24070122187", "CI/CD Action, SARIF & GitHub UI"], False, True)
    add_row(auth_table, w_auth, ["Shashank Tai", "24070122181", "Static Analysis, Rules & Secret Engine"], True, True)
    add_row(auth_table, w_auth, ["Sayantan Nandi", "24070122176", "Hybrid AI Engine, Redaction & Git Patch"], False, True)

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(12)
    p_sp.paragraph_format.space_after = Pt(12)

    # Project Overview Box (plain academic style)
    add_p(
        doc,
        "Abstract: SecureCode is an automated application security platform designed to address the high hosting costs, "
        "complex configuration, and developer burnout associated with traditional static analysis tools. Built as a "
        "lightweight GitHub Actions composite action, SecureCode scans code repositories on every commit and pull request "
        "with zero dedicated server infrastructure. It unites Semgrep Abstract Syntax Tree (AST) static analysis, "
        "Gitleaks high-entropy secret detection, privacy-preserving credential redaction, and Groq Language Processing "
        "Unit (LPU) artificial intelligence. The platform delivers three actionable outputs: an OASIS SARIF v2.1.0 "
        "report uploaded directly to GitHub's native Security tab, an executive dashboard inside GitHub Actions Step Summary, "
        "and a downloadable, unified Git remediation patch ('remediation.patch') that developers can apply instantly "
        "using 'git apply'."
    )

    doc.add_page_break()

    # -------------------------------------------------------------
    # CHAPTER 1: INTRODUCTION AND PROBLEM BACKGROUND
    # -------------------------------------------------------------
    add_heading_1(doc, "1. Introduction and Problem Background")

    add_p(
        doc,
        "Modern software development moves at a rapid pace. Continuous Integration and Continuous Deployment (CI/CD) "
        "pipelines build, test, and release code multiple times a day. In this fast-moving environment, security testing "
        "can no longer be treated as an isolated audit performed by an external team right before a product launch. "
        "When security flaws are discovered late, fixing them requires expensive architecture rewrites, delays release schedules, "
        "and risks severe data breaches. This reality gave rise to the philosophy of 'Shift-Left Security'—the practice of "
        "integrating security checks directly into the earliest stages of software development."
    )

    add_heading_2(doc, "1.1 The Practical Problems with Existing Security Tools")
    add_p(
        doc,
        "While the benefits of Shift-Left security are widely acknowledged, the tools currently available to teams "
        "suffer from severe practical limitations that prevent widespread adoption by students, open-source projects, "
        "and resource-constrained development teams:"
    )

    add_bullet(
        doc,
        "Persistent Hosting and Infrastructure Costs",
        "Established enterprise static analysis tools, such as SonarQube, Veracode, and Checkmarx, require persistent "
        "dedicated servers, background worker processes, and relational databases. Maintaining this dedicated cloud "
        "infrastructure costs anywhere from $50 to hundreds of dollars each month, which is prohibitive for small teams and individual developers."
    )
    add_bullet(
        doc,
        "Cryptic Jargon and Developer Fatigue",
        "Traditional scanners output cryptic rule identifiers such as 'CWE-89' or 'B301' buried inside thousands of lines "
        "of raw terminal logs. Developers are presented with a list of alarms without any clear explanation of why the "
        "pattern is dangerous, what realistic attack vector an adversary could use, or how to write an effective fix."
    )
    add_bullet(
        doc,
        "Absence of Ready-to-Apply Remediations",
        "When a security alert is triggered, developers must spend substantial time reading external documentation "
        "and drafting safe replacements manually. Existing CI tools do not provide a downloadable, pre-computed Git patch "
        "that can be reviewed and applied with a single command."
    )
    add_bullet(
        doc,
        "Privacy and Secret Exposure Risks with Third-Party AI",
        "Many developers currently resort to copying vulnerable code snippets into public AI chatbots for explanations. "
        "However, code often contains hardcoded API keys, database credentials, or private proprietary logic. Sending raw "
        "code to third-party language models without prior sanitization leads to severe confidentiality breaches and compliance violations."
    )

    add_heading_2(doc, "1.2 The SecureCode Solution")
    add_p(
        doc,
        "SecureCode was conceived and built to eliminate each of these barriers. The project delivers an enterprise-grade "
        "security pipeline with zero server footprint, running entirely inside the ephemeral containers of GitHub Actions. "
        "It combines AST-based code analysis, secret detection, regex-based privacy redaction, and fast Groq LPU inference "
        "to give developers clear explanations and automated Git patches right where they already work."
    )

    # -------------------------------------------------------------
    # CHAPTER 2: SYSTEM ARCHITECTURE AND DATA FLOW
    # -------------------------------------------------------------
    add_heading_1(doc, "2. System Architecture and End-to-End Workflow")

    add_p(
        doc,
        "SecureCode is architected as an autonomous composite GitHub Action. When a developer pushes code or opens a pull request, "
        "GitHub Actions provisions a clean Ubuntu container, initializes the SecureCode engine, executes the analysis, and "
        "publishes the security artifacts. The complete architecture is illustrated in Figure 1 below."
    )

    add_image_figure(
        doc,
        "diagrams/system_architecture.png",
        "Figure 1: High-Level System Architecture of SecureCode showing trigger, runner engines, and deliverable artifacts."
    )

    add_heading_2(doc, "2.1 Pipeline Stages Walkthrough")
    add_p(
        doc,
        "The end-to-end execution of SecureCode progresses through six structured stages:"
    )

    pipe_table = doc.add_table(rows=7, cols=3)
    pipe_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    w_pipe = [1.2, 1.8, 3.5]
    style_table_header(pipe_table.rows[0], w_pipe, ["Stage", "Engine / Component", "Operational Responsibilities"])
    add_row(pipe_table, w_pipe, ["1. Ingestion", "GitHub Actions Runner", "Pulls caller repository, installs Python 3.11, and downloads standalone Gitleaks binary in under 15 seconds."], False, True)
    add_row(pipe_table, w_pipe, ["2. Static Scan", "Semgrep & Gitleaks", "Semgrep evaluates AST syntax rules; Gitleaks scans git history and files for leaked tokens."], True, True)
    add_row(pipe_table, w_pipe, ["3. Normalization", "Finding Normalizer", "Strips system paths, maps clean human titles, calculates SHA-256 fingerprints, tallies severities."], False, True)
    add_row(pipe_table, w_pipe, ["4. Sanitization", "Privacy Redaction", "Slices surrounding code context (+/- 12 lines) and redacts secrets via regex before LLM prompt assembly."], True, True)
    add_row(pipe_table, w_pipe, ["5. Remediation", "Hybrid AI Engine", "Queries Groq LPU (openai/gpt-oss-120b) for structured JSON diffs; cascades to offline knowledge base if needed."], False, True)
    add_row(pipe_table, w_pipe, ["6. Deliverables", "Artifact Exporter", "Generates Step Summary markdown dashboard, results.sarif (v2.1.0), and unified remediation.patch."], True, True)

    p_sp2 = doc.add_paragraph()
    p_sp2.paragraph_format.space_before = Pt(6)

    add_heading_2(doc, "2.2 The Vulnerability Remediation Lifecycle")
    add_p(
        doc,
        "Figure 2 illustrates the detailed lifecycle of a single finding from its initial identification in source code "
        "to its compilation into the final patch file."
    )

    add_image_figure(
        doc,
        "diagrams/data_flow_pipeline.png",
        "Figure 2: Vulnerability Analysis & AI Patch Generation Pipeline showing context extraction and diff compilation."
    )

    # -------------------------------------------------------------
    # CHAPTER 3: DETAILED TECHNICAL BREAKDOWN OF ENGINES
    # -------------------------------------------------------------
    add_heading_1(doc, "3. Technical Deep-Dive: Core Engineering Modules")

    add_p(
        doc,
        "To provide a complete understanding of how SecureCode works under the hood, this chapter breaks down each of the "
        "seven internal software engines that make up the platform."
    )

    add_heading_2(doc, "3.1 Engine 1: Semgrep AST Rule Engine")
    add_p(
        doc,
        "Unlike naive scanners that rely on simple regular expressions (which produce overwhelming false positives on comments "
        "or benign string mentions), SecureCode utilizes Semgrep's Abstract Syntax Tree (AST) engine. When source code is scanned, "
        "Semgrep parses the text into a structural syntax tree representing function calls, variable assignments, and data flows. "
        "This allows SecureCode to enforce semantic security checks, such as identifying when user-controlled input reaches an "
        "operating system execution sink (e.g., os.system) or an insecure deserializer (pickle.loads), regardless of how the code is formatted."
    )
    add_p(
        doc,
        "The engine operates in bundled mode by default, loading pre-compiled, verified YAML rule packs stored in backend/rules. "
        "To maximize scan speed and prevent noise, the Semgrep runner automatically excludes third-party dependency directories "
        "including node_modules, venv, .git, and cache directories."
    )

    add_heading_2(doc, "3.2 Engine 2: Gitleaks Secret Detection Engine")
    add_p(
        doc,
        "Hardcoded credentials, API keys, and database passwords committed to version control represent an immediate security threat. "
        "SecureCode incorporates Gitleaks as a specialized secret hunter. Gitleaks combines regex signature matching with "
        "Shannon entropy analysis. High-entropy strings (random alphanumeric strings typical of cryptographic keys, JWT tokens, "
        "and cloud access keys) are flagged immediately. The Gitleaks runner executes with process isolation (shell=False) "
        "and writes structured JSON reports that are parsed into unified Finding models."
    )

    add_heading_2(doc, "3.3 Engine 3: Unified Normalizer and Deduplication Pipeline")
    add_p(
        doc,
        "A critical engineering hurdle in multi-scanner architectures is reconciling the raw outputs of different tools. "
        "Raw Semgrep outputs contain machine-specific path prefixes that look like: "
        "'home.runner.work._actions.ShlokGarg22.SecureCode.main.backend.rules.python-command-injection'. "
        "If displayed directly, this clutters the interface and confuses developers."
    )
    add_p(
        doc,
        "SecureCode addresses this in normalizer.py through the clean_rule_id_and_title() function. It strips path prefixes "
        "and maps each check to a clean, standardized title such as 'Command Injection via Shell (os.system)'. Furthermore, "
        "the normalizer computes a stable SHA-256 fingerprint from the rule ID, normalized file path, and code snippet. "
        "This fingerprint allows SecureCode to deduplicate recurring findings across scans and maintain a clean audit history."
    )

    add_heading_2(doc, "3.4 Engine 4: Automated Privacy Redaction Shield")
    add_p(
        doc,
        "Before sending any code to an external language model for explanation, SecureCode runs the snippet through "
        "redact_secrets_enhanced(). This privacy filter uses regular expression patterns to identify sensitive data: "
        "Stripe keys, AWS access keys, GitHub tokens, Slack webhooks, JWT tokens, and database connection strings with passwords. "
        "These strings are replaced with safe placeholder tokens (e.g., [REDACTED_AWS_KEY]). This ensures that proprietary "
        "credentials are never transmitted over the network to cloud AI providers."
    )

    add_heading_2(doc, "3.5 Engine 5: Hybrid AI Remediation Engine with Multi-Model Fallback")
    add_p(
        doc,
        "SecureCode features a dual-engine architecture designed for high availability and intelligence:"
    )
    add_bullet(
        doc,
        "Surrounding Source Context Extraction",
        "Instead of sending only a disconnected 1-line snippet, SecureCode reads the target file and extracts 25 surrounding "
        "lines (+/- 12 lines). This provides the language model with essential context regarding variable definitions, "
        "function signatures, and imports, enabling it to write accurate, syntactically correct code replacements."
    )
    add_bullet(
        doc,
        "Live Groq LPU Inference (openai/gpt-oss-120b)",
        "When an API key is provided, SecureCode connects to Groq's high-speed Language Processing Unit (LPU) cloud. "
        "It uses the verified 'openai/gpt-oss-120b' model with a 131,072 token context window. At a low temperature setting (0.1), "
        "the model analyzes the code and returns structured JSON containing plain explanations, realistic attack scenarios, and git diffs in ~1.2 seconds."
    )
    add_bullet(
        doc,
        "Autonomous Multi-Model Fallback Cascade",
        "If a specific model is rate-limited (due to Groq TPM thresholds) or unavailable, the engine automatically attempts "
        "alternative candidate models in sequence: openai/gpt-oss-120b -> openai/gpt-oss-20b -> openai/gpt-oss-safeguard-20b -> allam-2-7b. "
        "This ensures that temporary cloud quotas never block a developer's CI build."
    )
    add_bullet(
        doc,
        "Offline Deterministic Fallback",
        "If no API key is configured or the runner is operating in an air-gapped environment, SecureCode falls back to its "
        "built-in deterministic security knowledge base. Pre-computed, verified remediations are generated locally without network access."
    )

    add_heading_2(doc, "3.6 Engine 6: OASIS SARIF v2.1.0 Exporter")
    add_p(
        doc,
        "The Static Analysis Results Interchange Format (SARIF) is the official OASIS standard for security diagnostics. "
        "SecureCode's sarif_exporter.py maps internal ScanResult models into SARIF v2.1.0 JSON. Crucially, code replacements "
        "are wrapped strictly under 'artifactChanges' as required by the SARIF specification. This allows GitHub's CodeQL "
        "action to ingest the file and populate native security alerts inside the repository's 'Security' tab (Figure 3)."
    )

    add_image_figure(
        doc,
        "diagrams/github_sarif_integration.png",
        "Figure 3: OASIS SARIF v2.1.0 GitHub Security Tab Integration mapping scan results to native code scanning alerts."
    )

    add_heading_2(doc, "3.7 Engine 7: Unified Git Patch Generator")
    add_p(
        doc,
        "The final module in the pipeline is patch_generator.py. It collects all validated remediation diffs across all files "
        "and formats them into a standard unified Git patch (remediation.patch). The patch includes complete file headers "
        "('--- a/file', '+++ b/file') and chunk coordinates ('@@ -start,len +start,len @@'). A developer can download this "
        "artifact and run 'git apply remediation.patch' to resolve all detected vulnerabilities simultaneously."
    )

    doc.add_page_break()

    # -------------------------------------------------------------
    # CHAPTER 4: TEAM CONTRIBUTIONS AND WORK DISTRIBUTION
    # -------------------------------------------------------------
    add_heading_1(doc, "4. Team Member Contributions and Work Distribution")

    add_p(
        doc,
        "SecureCode was designed, built, and tested through equal collaboration among three computer engineering students. "
        "To ensure clean modularity and accountability, each team member took primary ownership of specific subsystems "
        "spanning CI/CD infrastructure, static analysis engines, and AI remediation services."
    )

    # 4.1 Shlok Garg
    add_heading_2(doc, "4.1 Shlok Garg (PRN: 24070122187)")
    add_p(doc, "Primary Ownership: CI/CD Composite Action Architecture, SARIF v2.1.0 Engine & GitHub UI Dashboard")
    add_bullet(
        doc,
        "Composite GitHub Action (action.yml)",
        "Designed and implemented the standalone GitHub Action definition. Configured environment setup, Python 3.11 runtime, "
        "dependency installation, and automated binary compilation for Gitleaks."
    )
    add_bullet(
        doc,
        "OASIS SARIF v2.1.0 Exporter (sarif_exporter.py)",
        "Engineered the SARIF generation module. Identified and resolved a critical GitHub CodeQL schema rejection by "
        "implementing the 'artifactChanges' structure, enabling seamless ingestion into GitHub's native Security tab."
    )
    add_bullet(
        doc,
        "CI Policy and Return-Code Trapping (scan.py)",
        "Engineered the threshold policy evaluation logic (fail-on: critical, high, medium). Implemented bash return-code "
        "trapping in action.yml so SARIF reports and patches are guaranteed to write before runner termination."
    )
    add_bullet(
        doc,
        "Step Summary Markdown Dashboard (format_github_step_summary)",
        "Designed the executive dashboard rendered in GitHub Actions. Structured clean severity breakdown tables, auto-expanded "
        "Critical/High finding cards, and multi-line diff blocks."
    )
    add_bullet(
        doc,
        "Integration & Testbed Deployment",
        "Set up the public target demo repository (ShlokGarg22/TestingSecureCode) and validated continuous workflow execution."
    )

    # 4.2 Shashank Tai
    add_heading_2(doc, "4.2 Shashank Tai (PRN: 24070122181)")
    add_p(doc, "Primary Ownership: Static Analysis Engines, Security Rule Engineering & Secret Detection")
    add_bullet(
        doc,
        "Semgrep AST Runner Engine (semgrep.py)",
        "Engineered the process execution wrapper for Semgrep, managing rule mode selection (bundled, online, hybrid), "
        "process timeouts, and smart directory exclusion (skipping vendor dependencies and cache folders)."
    )
    add_bullet(
        doc,
        "Gitleaks Secret Detection Runner (gitleaks.py)",
        "Integrated the Gitleaks binary, configuring command execution with process isolation (shell=False) and parsing "
        "raw JSON reports into structured secret finding models."
    )
    add_bullet(
        doc,
        "Security Rule Curation (backend/rules)",
        "Authored and tuned AST pattern rules targeting critical Python vulnerabilities, including OS command injection, "
        "pickle deserialization, path traversal, and dangerous eval() expressions."
    )
    add_bullet(
        doc,
        "Finding Normalization Pipeline (normalizer.py)",
        "Created the normalizer that standardizes tool outputs. Authored clean_rule_id_and_title() to strip path clutter "
        "and implemented SHA-256 content fingerprinting to eliminate duplicate alerts across runs."
    )
    add_bullet(
        doc,
        "Severity Classification Algorithm",
        "Formulated the severity weighting matrix (Critical, High, Medium, Low) mapping CWE risk categories to actionable CI thresholds."
    )

    # 4.3 Sayantan Nandi
    add_heading_2(doc, "4.3 Sayantan Nandi (PRN: 24070122176)")
    add_p(doc, "Primary Ownership: Hybrid AI Remediation Engine, Privacy Redaction & Unified Git Patch")
    add_bullet(
        doc,
        "Hybrid AI Explainer Architecture (explainer_enhanced.py)",
        "Architected the dual AI engine integrating Groq LPU inference (openai/gpt-oss-120b, openai/gpt-oss-20b) with "
        "an offline deterministic knowledge base fallback."
    )
    add_bullet(
        doc,
        "Multi-Model Fallback Cascade",
        "Engineered the fault-tolerant model cascade that dynamically shifts across candidate Groq models upon encountering "
        "TPM rate limits, ensuring continuous CI execution."
    )
    add_bullet(
        doc,
        "Privacy & Secret Redaction Pipeline (redact_secrets_enhanced)",
        "Designed the regex sanitization filter that intercepts and redacts high-entropy API keys, AWS credentials, and "
        "tokens before prompt construction."
    )
    add_bullet(
        doc,
        "Unified Git Patch Generator (patch_generator.py)",
        "Engineered the multi-diff patch compiler that parses AI remediation outputs and formats them into standard unified "
        "diffs (remediation.patch) compatible with git apply."
    )
    add_bullet(
        doc,
        "Groq Model Diagnostic Tool (check_groq_models.py)",
        "Developed the standalone CLI utility that queries Groq client.models.list(), benchmarks latency, and validates JSON schemas."
    )
    add_bullet(
        doc,
        "Full-Stack Web Interface & API Integration",
        "Co-developed the FastAPI backend endpoints and Next.js web application for manual ZIP uploads and side-by-side diff previews."
    )

    # Summary Matrix Table
    add_heading_2(doc, "4.4 Responsibility and Deliverables Summary")
    resp_table = doc.add_table(rows=6, cols=4)
    resp_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    w_resp = [2.2, 1.4, 1.4, 1.5]
    style_table_header(resp_table.rows[0], w_resp, ["Module / Subsystem", "Lead Owner", "Secondary", "Deliverable Files"])
    add_row(resp_table, w_resp, ["CI/CD Composite Action", "Shlok Garg", "Sayantan Nandi", "action.yml, securecode.yml"], False, True)
    add_row(resp_table, w_resp, ["Static Analysis & Rules", "Shashank Tai", "Shlok Garg", "semgrep.py, gitleaks.py, rules/"], True, True)
    add_row(resp_table, w_resp, ["SARIF v2.1.0 & Step Summary", "Shlok Garg", "Shashank Tai", "sarif_exporter.py, scan.py"], False, True)
    add_row(resp_table, w_resp, ["Hybrid AI Engine & Fallback", "Sayantan Nandi", "Shashank Tai", "explainer_enhanced.py"], True, True)
    add_row(resp_table, w_resp, ["Git Patch & Privacy Shield", "Sayantan Nandi", "Shlok Garg", "patch_generator.py, check_groq_models.py"], False, True)

    doc.add_page_break()

    # -------------------------------------------------------------
    # CHAPTER 5: REAL-WORLD TESTING AND EVALUATION
    # -------------------------------------------------------------
    add_heading_1(doc, "5. Real-World Testing and Demonstration Results")

    add_p(
        doc,
        "To validate SecureCode in a realistic production scenario, we created a test repository named 'TestingSecureCode' "
        "containing representative code vulnerabilities in a Python application (app.py). The repository was configured "
        "with GitHub Actions to run SecureCode on every commit."
    )

    add_heading_2(doc, "5.1 Evaluated Vulnerabilities in app.py")
    eval_table = doc.add_table(rows=7, cols=4)
    eval_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    w_eval = [0.6, 1.2, 2.2, 2.5]
    style_table_header(eval_table.rows[0], w_eval, ["#", "Severity", "Vulnerability Class", "Vulnerable Code & Mechanism"])
    add_row(eval_table, w_eval, ["1", "Critical", "Command Injection (os.system)", "app.py:117 — os.system(f'ping -c 1 {ip}') executes unsanitized shell commands."], False, True)
    add_row(eval_table, w_eval, ["2", "Critical", "Insecure Pickle Deserialization", "app.py:151 — pickle.loads(raw_bytes) allows arbitrary bytecode execution."], True, True)
    add_row(eval_table, w_eval, ["3", "High", "Dangerous eval() Execution", "app.py:165 — computed = eval(expr) dynamically executes untrusted string formulas."], False, True)
    add_row(eval_table, w_eval, ["4", "High", "Hardcoded API Secret Credential", "app.py:18 — AWS_ACCESS_KEY committed directly in source code."], True, True)
    add_row(eval_table, w_eval, ["5", "Medium", "Path Traversal (open)", "app.py:30 — open(f'uploads/{user_path}') allows ../ directory traversal."], False, True)
    add_row(eval_table, w_eval, ["6", "Medium", "Path Traversal (open)", "app.py:106 — open(doc_path) accesses files outside authorized base root."], True, True)

    p_sp3 = doc.add_paragraph()
    p_sp3.paragraph_format.space_before = Pt(6)

    add_heading_2(doc, "5.2 Live Scan Output: GitHub Actions Step Summary")
    add_p(
        doc,
        "When triggered on GitHub Actions, the workflow ran in 27 seconds. SecureCode parsed all six findings and formatted "
        "the following executive review directly inside the Actions Summary page:"
    )

    add_code(
        doc,
        "SecureCode Security Audit Report\n"
        "Project: TestingSecureCode | Scan ID: CLI-20261021123255 | Status: Policy Blocked (fail-on: high)\n\n"
        "Security Posture Matrix\n"
        "---------------------------------------------------------------------------\n"
        "Critical: 2 findings (Remote Code Execution / Host Compromise) - [Blocker]\n"
        "High:     2 findings (Direct Exploitation / Exposed Secrets)   - [Blocker]\n"
        "Medium:   2 findings (Path Traversal / Injection Vectors)      - [Review]\n"
        "Low:      0 findings (Configuration Hardening)                 - [Clean]\n"
        "---------------------------------------------------------------------------\n\n"
        "1. [Critical] Command Injection via Shell (os.system) — app.py:117\n"
        "   - Location: app.py (Line 117)\n"
        "   - Analysis Engine: Live AI (Groq gpt-oss-120b)\n"
        "   - Why It Matters: User input is passed directly to the operating system shell.\n"
        "   - Attack Scenario: Attacker passes '127.0.0.1; cat /etc/passwd' to exfiltrate system accounts.\n"
        "   - Suggested Remediation Diff:\n"
        "     - command = f'ping -c 1 {target_ip}'\n"
        "     - os.system(command)\n"
        "     + subprocess.run(['ping', '-c', '1', target_ip], capture_output=True, shell=False)",
        caption="Sample Output Generated in GitHub Actions Step Summary"
    )

    add_heading_2(doc, "5.3 Applying the Generated Git Patch")
    add_p(
        doc,
        "All six AI-generated diffs were compiled into a single downloadable artifact named 'remediation-patch'. "
        "A developer applies all fixes locally using standard Git commands:"
    )

    add_code(
        doc,
        "# In the repository root directory:\n"
        "$ git apply remediation.patch\n"
        "$ git status\n"
        "modified:   app.py\n"
        "$ git commit -am 'fix: apply automated security remediation patch from SecureCode'\n"
        "$ git push origin main",
        caption="Terminal Workflow to Apply Remediation Patch"
    )

    # -------------------------------------------------------------
    # CHAPTER 6: COMPARATIVE EVALUATION AND PERFORMANCE
    # -------------------------------------------------------------
    add_heading_1(doc, "6. Comparative Evaluation and Benchmarking")

    add_p(
        doc,
        "To evaluate SecureCode objectively, we compared its features, operational cost, and developer experience "
        "against three leading industry tools: SonarQube, Snyk, and GitHub CodeQL."
    )

    comp_table = doc.add_table(rows=6, cols=5)
    comp_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    w_comp = [1.7, 1.2, 1.2, 1.2, 1.2]
    style_table_header(comp_table.rows[0], w_comp, ["Evaluation Metric", "SecureCode", "SonarQube", "Snyk", "CodeQL"])
    add_row(comp_table, w_comp, ["Hosting Cost", "$0 (Zero-Hosting)", "$50 - $200/month", "SaaS Subscription", "$0 (Public Repos)"], False, True)
    add_row(comp_table, w_comp, ["Setup Effort", "1 YAML file (10 lines)", "Complex (DB + Server)", "CLI / SaaS Token", "Complex Build Config"], True, True)
    add_row(comp_table, w_comp, ["AI Attack Scenarios", "Yes (Groq LPU)", "No", "Basic Advice", "No"], False, True)
    add_row(comp_table, w_comp, ["Downloadable Patch", "Yes (remediation.patch)", "No", "PR Bot Only", "No"], True, True)
    add_row(comp_table, w_comp, ["Average Runtime", "~25 - 30 seconds", "~2 - 5 minutes", "~1 - 2 minutes", "~3 - 8 minutes"], False, True)

    p_sp4 = doc.add_paragraph()
    p_sp4.paragraph_format.space_before = Pt(6)

    add_heading_2(doc, "6.1 Performance and Latency Breakdown")
    add_bullet(doc, "Gitleaks Binary Execution", "1.1 seconds — Fast compiled Go binary scanning local git buffers.")
    add_bullet(doc, "Semgrep AST Static Analysis", "4.2 seconds — Efficient rule evaluation across code syntax trees.")
    add_bullet(doc, "Groq LPU AI Inference", "1.2 seconds per finding — Ultra-fast inference powered by Groq LPUs.")
    add_bullet(doc, "Artifact Compilation", "0.4 seconds — Assembling SARIF report and remediation.patch.")
    add_bullet(doc, "Total Runner Execution", "27.0 seconds total time from workflow trigger to completion.")

    # -------------------------------------------------------------
    # CHAPTER 7: ENGINEERING CHALLENGES AND SOLUTIONS
    # -------------------------------------------------------------
    add_heading_1(doc, "7. Engineering Challenges and Practical Solutions")

    add_heading_2(doc, "7.1 Challenge 1: Python Module Resolution in Composite Actions")
    add_p(
        doc,
        "When executing a composite GitHub Action, the working directory is the caller's repository, not the action's "
        "source repository. Initial runs failed with 'ModuleNotFoundError: No module named app'. We resolved this by "
        "dynamically prepending the action's backend directory to sys.path at the very top of scan.py: "
        "backend_dir = Path(__file__).resolve().parent.parent; sys.path.insert(0, str(backend_dir)). "
        "This ensured that internal imports resolve reliably across any repository."
    )

    add_heading_2(doc, "7.2 Challenge 2: OASIS SARIF Schema Incompatibility")
    add_p(
        doc,
        "GitHub's CodeQL action performs strict schema validation on uploaded SARIF files. Initial runs failed with: "
        "'fixes[0] is not allowed to have additional property fileChanges; requires artifactChanges'. "
        "Research revealed that the OASIS SARIF v2.1.0 standard strictly designates 'artifactChanges' as the container "
        "for file replacements. We refactored sarif_exporter.py accordingly, achieving full compliance and enabling "
        "native alerts in GitHub's Security tab."
    )

    add_heading_2(doc, "7.3 Challenge 3: Preserving Multi-Line Git Diffs in Markdown")
    add_p(
        doc,
        "In GitHub Actions Step Summary, omitting blank lines immediately after HTML <details> and <summary> tags caused "
        "GitHub's CommonMark parser to collapse entire multi-line diff blocks into a single illegible line. We revised "
        "format_github_step_summary() to enforce proper whitespace separation and convert escaped newlines, ensuring "
        "clean multi-line code diff rendering."
    )

    add_heading_2(doc, "7.4 Challenge 4: Groq Model Availability and Account Tiers")
    add_p(
        doc,
        "Initial configurations attempted to call Meta LLaMA 3.3 models on Groq. However, different Groq accounts "
        "and geographical regions have varying model permissions, causing access and token-per-minute (TPM) errors. "
        "We built check_groq_models.py to programmatically query client.models.list(), benchmarking and discovering "
        "that 'openai/gpt-oss-120b' and 'openai/gpt-oss-20b' were active and provided 131,072 token context windows. "
        "We reconfigured the explainer to prioritize these models, backed by an autonomous multi-model fallback cascade."
    )

    # -------------------------------------------------------------
    # CHAPTER 8: CONCLUSION AND FUTURE ROADMAP
    # -------------------------------------------------------------
    add_heading_1(doc, "8. Conclusion and Future Roadmap")

    add_p(
        doc,
        "SecureCode demonstrates that robust DevSecOps security does not require costly cloud servers, complex database setups, "
        "or frustrating developer workflows. By uniting AST static analysis, high-entropy secret detection, privacy-preserving "
        "redaction, and fast Groq LPU inference, SecureCode delivers automated security reviews and actionable Git patches "
        "directly inside GitHub Actions with zero hosting costs."
    )

    add_heading_2(doc, "8.1 Future Roadmap")
    add_bullet(doc, "Multi-Language Expansion", "Extend bundled AST rules to JavaScript, TypeScript, Go, Java, and Rust.")
    add_bullet(doc, "Automated Pull Request Bot", "Build a bot to automatically create branch pull requests containing the applied remediation patch, allowing developers to merge fixes with one click.")
    add_bullet(doc, "IDE Extension", "Package the deterministic and AI remediation engines into a VS Code extension for inline security warnings before code is committed.")

    # Save Document
    out_path = os.path.abspath("SecureCode_Project_Report.docx")
    doc.save(out_path)
    print(f"\n[+] Successfully generated SecureCode Project Report at:\n    {out_path}")
    return out_path


if __name__ == "__main__":
    build_report()
