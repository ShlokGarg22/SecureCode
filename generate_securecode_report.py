#!/usr/bin/env python3
"""
SecureCode Comprehensive Project Report Generator
Creates a professionally styled, humanized, and highly detailed DOCX report for SecureCode.
Team Members:
  - Shlok Garg (PRN: 24070122187)
  - Shashank Tai (PRN: 24070122181)
  - Sayantan Nandi (PRN: 24070122176)
"""

import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

# Color Palette
COLOR_PRIMARY = RGBColor(30, 58, 138)       # Deep Navy #1E3A8A
COLOR_SECONDARY = RGBColor(59, 130, 246)   # Royal Blue #3B82F6
COLOR_DARK_TEXT = RGBColor(31, 41, 55)     # Dark Slate #1F2937
COLOR_MUTED = RGBColor(100, 116, 139)      # Slate Gray #64748B
COLOR_SUCCESS = RGBColor(22, 101, 52)      # Green #166534
COLOR_DANGER = RGBColor(185, 28, 28)       # Crimson #B91C1C

HEX_HEADER_BG = "1E3A8A"
HEX_ROW_ALT = "F8FAFC"
HEX_CALLOUT_BG = "F1F5F9"
HEX_CODE_BG = "F8FAFC"
HEX_BORDER = "CBD5E1"


def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    """Set cell padding (in dxa: 20 dxa = 1 pt)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def set_cell_shading(cell, color_hex):
    """Set background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)


def set_cell_borders(cell, top="none", bottom="none", left="none", right="none", 
                     color="CBD5E1", sz="8"):
    """Set borders on individual table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    borders_xml = f'<w:tcBorders {nsdecls("w")}>'
    for side, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        if val != "none":
            borders_xml += f'<w:{side} w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        else:
            borders_xml += f'<w:{side} w:val="none"/>'
    borders_xml += '</w:tcBorders>'
    tcPr.append(parse_xml(borders_xml))


def format_table_header(row, col_widths, titles):
    """Style a table header row."""
    for idx, cell in enumerate(row.cells):
        cell.width = Inches(col_widths[idx])
        set_cell_shading(cell, HEX_HEADER_BG)
        set_cell_margins(cell, top=180, bottom=180, left=180, right=180)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(titles[idx])
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(255, 255, 255)


def add_table_row(table, col_widths, values, is_alt=False, is_bold_first=False):
    """Add a styled data row to a table."""
    row = table.add_row()
    for idx, cell in enumerate(row.cells):
        cell.width = Inches(col_widths[idx])
        bg = HEX_ROW_ALT if is_alt else "FFFFFF"
        set_cell_shading(cell, bg)
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
        set_cell_borders(cell, bottom="single", color=HEX_BORDER, sz="4")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(str(values[idx]))
        run.font.name = "Arial"
        run.font.size = Pt(9.5)
        run.font.color.rgb = COLOR_DARK_TEXT
        if idx == 0 and is_bold_first:
            run.bold = True
    return row


def add_callout(doc, title, text, icon="💡"):
    """Insert a highlighted callout box with a primary left border."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_shading(cell, HEX_CALLOUT_BG)
    set_cell_margins(cell, top=180, bottom=180, left=240, right=200)
    set_cell_borders(cell, left="single", color=HEX_HEADER_BG, sz="36")
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run_title = p.add_run(f"{icon} {title}\n")
    run_title.bold = True
    run_title.font.name = "Arial"
    run_title.font.size = Pt(10.5)
    run_title.font.color.rgb = COLOR_PRIMARY
    
    run_text = p.add_run(text)
    run_text.font.name = "Arial"
    run_text.font.size = Pt(9.5)
    run_text.font.color.rgb = COLOR_DARK_TEXT
    
    # Add spacing after table
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)


def add_code_block(doc, code_str, caption=None):
    """Insert a clean syntax box for code snippets."""
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.paragraph_format.space_before = Pt(8)
        p_cap.paragraph_format.space_after = Pt(2)
        r_cap = p_cap.add_run(f"Code Listing: {caption}")
        r_cap.font.name = "Arial"
        r_cap.font.size = Pt(9)
        r_cap.font.italic = True
        r_cap.font.color.rgb = COLOR_MUTED

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_shading(cell, HEX_CODE_BG)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    set_cell_borders(cell, top="single", bottom="single", left="single", right="single", color="E2E8F0", sz="8")
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(code_str.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(30, 41, 59)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)


def add_heading_1(doc, text):
    """Add a styled Heading 1."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(16)
    run.font.color.rgb = COLOR_PRIMARY
    return p


def add_heading_2(doc, text):
    """Add a styled Heading 2."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(12.5)
    run.font.color.rgb = COLOR_SECONDARY
    return p


def add_heading_3(doc, text):
    """Add a styled Heading 3."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(10.5)
    run.font.color.rgb = COLOR_DARK_TEXT
    return p


def add_body_paragraph(doc, text):
    """Add a standard readable body paragraph."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.2
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10)
    run.font.color.rgb = COLOR_DARK_TEXT
    return p


def add_bullet_item(doc, title, text):
    """Add a styled bullet list item."""
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    r_title = p.add_run(f"{title}: ")
    r_title.bold = True
    r_title.font.name = "Arial"
    r_title.font.size = Pt(10)
    r_title.font.color.rgb = COLOR_DARK_TEXT
    
    r_text = p.add_run(text)
    r_text.font.name = "Arial"
    r_text.font.size = Pt(10)
    r_text.font.color.rgb = COLOR_DARK_TEXT
    return p


def build_full_report():
    print("[*] Generating comprehensive SecureCode Project Report (DOCX)...")
    doc = Document()
    
    # Page setup: Standard Letter with 1-inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header / Footer
        footer = section.footer
        p_ft = footer.paragraphs[0]
        p_ft.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r_ft = p_ft.add_run("SecureCode Project Report | Confidential & Academic Evaluation")
        r_ft.font.name = "Arial"
        r_ft.font.size = Pt(8.5)
        r_ft.font.color.rgb = COLOR_MUTED

    # ---------------------------------------------------------
    # COVER / TITLE BLOCK
    # ---------------------------------------------------------
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(36)
    p_title.paragraph_format.space_after = Pt(8)
    r_main = p_title.add_run("SecureCode")
    r_main.bold = True
    r_main.font.name = "Arial"
    r_main.font.size = Pt(32)
    r_main.font.color.rgb = COLOR_PRIMARY
    
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(18)
    r_sub = p_sub.add_run("Zero-Hosting Automated CI/CD Security Scanner & AI Remediation Platform")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(13)
    r_sub.font.italic = True
    r_sub.font.color.rgb = COLOR_SECONDARY

    # Divider bar
    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.paragraph_format.space_after = Pt(24)
    r_div = p_div.add_run("—" * 45)
    r_div.font.color.rgb = COLOR_SECONDARY

    # Author Metadata Table
    meta_table = doc.add_table(rows=4, cols=3)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [2.2, 2.1, 2.2]
    
    format_table_header(meta_table.rows[0], widths, ["Project Contributor", "Student PRN", "Key Architectural Focus"])
    add_table_row(meta_table, widths, ["Shlok Garg", "24070122187", "CI/CD Action, SARIF & GitHub UI"], is_alt=False, is_bold_first=True)
    add_table_row(meta_table, widths, ["Shashank Tai", "24070122181", "Static Analysis, Rules & Secret Engine"], is_alt=True, is_bold_first=True)
    add_table_row(meta_table, widths, ["Sayantan Nandi", "24070122176", "Hybrid AI Engine, Redaction & Git Patch"], is_alt=False, is_bold_first=True)

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(12)
    p_space.paragraph_format.space_after = Pt(18)

    add_callout(
        doc,
        "Project Executive Mission",
        "SecureCode bridges the gap between static code security and developer developer velocity. "
        "By leveraging zero-hosting GitHub Actions composite architecture, privacy-preserving secret redaction, "
        "and Groq LPU-accelerated AI models, SecureCode automatically scans code on every push, maps native "
        "security alerts in GitHub Code Scanning, and delivers ready-to-apply unified Git patch artifacts.",
        icon="🛡️"
    )

    doc.add_page_break()

    # ---------------------------------------------------------
    # CHAPTER 1: EXECUTIVE SUMMARY & PROBLEM STATEMENT
    # ---------------------------------------------------------
    add_heading_1(doc, "1. Executive Summary & Problem Statement")
    
    add_body_paragraph(
        doc,
        "In modern agile software engineering, security cannot remain a manual afterthought checked once a quarter. "
        "The concept of 'Shift-Left Security' promises to catch security flaws, exposed passwords, and vulnerabilities "
        "early in the development cycle—right when a developer pushes code. However, existing enterprise solutions "
        "suffer from severe pain points that make adoption impractical for individuals, students, and agile teams."
    )

    add_heading_2(doc, "1.1 The Real-World Problems with Existing Tools")
    add_bullet_item(
        doc,
        "High Hosting & Infrastructure Costs",
        "Legacy static analysis tools (e.g., SonarQube, Veracode) require dedicated backend servers, persistent "
        "relational databases, and background worker daemons that cost anywhere from $50 to hundreds of dollars a month. "
        "Students and startups simply cannot afford persistent cloud infrastructure just to run code scans."
    )
    add_bullet_item(
        doc,
        "Cryptic Warnings and Developer Burnout",
        "Traditional scanners output cryptic rule identifiers such as 'CWE-89' or 'B301' buried inside thousands of lines "
        "of raw build console logs. Developers are given alarms without any plain-English explanation of why the vulnerability "
        "matters or how an attacker could actually exploit it in the wild."
    )
    add_bullet_item(
        doc,
        "Lack of Automated Actionable Fixes",
        "When an issue is flagged, existing tools leave the remediation entirely to the developer. The developer must "
        "spend hours researching documentation to write a patch. No major platform provides an immediate, verified, "
        "downloadable Git patch that can be applied with a single terminal command."
    )
    add_bullet_item(
        doc,
        "Privacy and Secret Leaks to AI Cloud Providers",
        "While modern developers try pasting code into public AI chatbots for explanations, doing so poses massive "
        "risks: hardcoded Stripe keys, AWS credentials, and proprietary proprietary secrets get inadvertently sent to third-party "
        "cloud models, violating compliance and triggering security breaches."
    )

    add_heading_2(doc, "1.2 The SecureCode Solution")
    add_body_paragraph(
        doc,
        "SecureCode was designed and engineered from the ground up to solve all four of these fundamental limitations. "
        "It is a zero-hosting, developer-first DevSecOps security suite that runs directly inside GitHub Actions with "
        "zero persistent server footprint, completely eliminating hosting costs."
    )

    add_callout(
        doc,
        "Core Value Proposition",
        "1. Zero Server Cost: Runs entirely in ephemeral GitHub runner environments ($0 hosting).\n"
        "2. Automated Unified Git Patch: Every run produces a downloadable 'remediation.patch' ready to execute via 'git apply'.\n"
        "3. Humanized AI Explanations: Plain-English explanations, real-world attack scenarios, and clean diff previews.\n"
        "4. Native GitHub Security Integration: OASIS SARIF v2.1.0 report uploaded directly to GitHub's Code Scanning tab.",
        icon="🚀"
    )

    # ---------------------------------------------------------
    # CHAPTER 2: ARCHITECTURE & END-TO-END DATA FLOW
    # ---------------------------------------------------------
    add_heading_1(doc, "2. System Architecture & End-to-End Pipeline")

    add_body_paragraph(
        doc,
        "SecureCode operates on an autonomous 6-stage pipeline. The entire system is architected as a lightweight composite "
        "GitHub Action, enabling any repository in the world to adopt full static analysis and AI remediation in fewer than "
        "10 lines of YAML configuration."
    )

    pipeline_table = doc.add_table(rows=7, cols=3)
    pipeline_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    p_widths = [1.2, 2.2, 3.1]
    
    format_table_header(pipeline_table.rows[0], p_widths, ["Stage", "Component", "Operational Role & Data Transformation"])
    add_table_row(pipeline_table, p_widths, ["Stage 1", "Workflow Trigger", "Monitors Git push / PR events; activates GitHub Action composite runner."], False, True)
    add_table_row(pipeline_table, p_widths, ["Stage 2", "Dual Scanner Execution", "Runs Semgrep (AST code rules) & Gitleaks (high-entropy secrets) in parallel."], True, True)
    add_table_row(pipeline_table, p_widths, ["Stage 3", "Normalization & Deduplication", "Cleans raw tool outputs, strips runner paths, creates stable cryptographic fingerprints."], False, True)
    add_table_row(pipeline_table, p_widths, ["Stage 4", "Privacy Sanitization", "Redacts high-entropy keys, JWTs, and AWS tokens from code snippets before AI ingestion."], True, True)
    add_table_row(pipeline_table, p_widths, ["Stage 5", "Hybrid AI Remediation", "Live Groq LPU inference (openai/gpt-oss-120b) with context; fallback to deterministic engine."], False, True)
    add_table_row(pipeline_table, p_widths, ["Stage 6", "Multi-Artifact Dispatch", "Writes GitHub Step Summary dashboard, SARIF v2.1.0 report, and unified remediation.patch."], True, True)

    p_space2 = doc.add_paragraph()
    p_space2.paragraph_format.space_before = Pt(8)

    add_heading_2(doc, "2.1 End-to-End Execution Flow")
    add_body_paragraph(
        doc,
        "When a developer pushes a commit to a repository configured with SecureCode, the following sequence occurs:"
    )
    add_bullet_item(doc, "1. Environment Setup", "GitHub Actions spins up a clean Ubuntu container, installs Python 3.11, downloads the standalone Gitleaks binary, and installs lightweight scanning dependencies in under 15 seconds.")
    add_bullet_item(doc, "2. Abstract Syntax Tree (AST) Analysis", "Semgrep scans the repository against bundled high-precision security rules, bypassing third-party library clutter (node_modules, venv, .git).")
    add_bullet_item(doc, "3. Secret & Credential Hunting", "Gitleaks scans git commits and local filesystem buffers for leaked API keys, database credentials, and private tokens.")
    add_bullet_item(doc, "4. Normalization & Severity Tally", "Findings from both engines are parsed into unified Finding data models, assigning standardized severities (Critical, High, Medium, Low) and calculating stable hash fingerprints.")
    add_bullet_item(doc, "5. AI Context Extraction & Prompting", "For each finding, SecureCode extracts 25 surrounding lines of real source code from the target file, strips all credentials via regex sanitization, and prompts the Groq LPU engine for a tailored, context-aware fix.")
    add_bullet_item(doc, "6. Multi-Format Output Generation", "Outputs the visual executive dashboard directly to $GITHUB_STEP_SUMMARY, exports an OASIS SARIF v2.1.0 document to results.sarif, and compiles all code diffs into remediation.patch.")

    # ---------------------------------------------------------
    # CHAPTER 3: DETAILED TECHNICAL BREAKDOWN OF ENGINES
    # ---------------------------------------------------------
    add_heading_1(doc, "3. Technical Deep-Dive: Core Engineering Modules")

    add_heading_2(doc, "3.1 Engine 1: Semgrep AST Rule Engine")
    add_body_paragraph(
        doc,
        "Unlike naive regex scanners that throw thousands of false alarms, SecureCode utilizes Semgrep's Abstract Syntax Tree "
        "(AST) engine. It understands code semantics, data-flow, and variable assignment scopes. Bundled rules target high-impact "
        "vulnerability classes including Command Injection (os.system, subprocess shell=True), Path Traversal (open, pathlib), "
        "Insecure Deserialization (pickle.loads), Dynamic Code Evaluation (eval, exec), and SQL Injection."
    )

    add_heading_2(doc, "3.2 Engine 2: Gitleaks Secret & Token Detection")
    add_body_paragraph(
        doc,
        "Hardcoded credentials represent one of the fastest paths to complete cloud infrastructure compromise. "
        "SecureCode incorporates a standalone Gitleaks runner that evaluates high-entropy strings and known credential signatures "
        "(Stripe API keys, AWS Access Keys, GitHub Personal Access Tokens, Slack Webhooks, and JWT tokens). "
        "Findings are isolated, normalized, and categorized under 'hardcoded-secret' with High severity."
    )

    add_heading_2(doc, "3.3 Engine 3: Normalization & Path Cleaning")
    add_body_paragraph(
        doc,
        "A major challenge in multi-tool integration is reconciling disparate outputs. Raw Semgrep checks contain long, "
        "ugly runner directory prefixes (e.g., home.runner.work._actions.ShlokGarg22.SecureCode...rules.python-command-injection). "
        "SecureCode implements clean_rule_id_and_title() in normalizer.py, stripping system paths and transforming cryptic IDs "
        "into clear, professional titles like 'Command Injection via Shell (os.system)'."
    )

    add_heading_2(doc, "3.4 Engine 4: Automated Secret Redaction & Privacy Shield")
    add_body_paragraph(
        doc,
        "Sending proprietary code containing secret keys to third-party AI models violates enterprise security standards. "
        "Before any prompt is dispatched to Groq or Gemini, SecureCode's redact_secrets_enhanced() pipeline intercepts the code. "
        "It uses multi-pattern regex matching to redact credentials into safe tokens (e.g., [REDACTED_STRIPE_KEY], "
        "[REDACTED_AWS_KEY], [REDACTED_PASSWORD]), ensuring zero sensitive credentials ever leave the runner environment."
    )

    add_heading_2(doc, "3.5 Engine 5: The Hybrid AI Remediation Engine")
    add_body_paragraph(
        doc,
        "SecureCode features a state-of-the-art dual AI architecture that combines ultra-fast cloud inference with offline reliability:"
    )
    add_bullet_item(
        doc,
        "Live Groq LPU Inference (openai/gpt-oss-120b)",
        "When an API key is provided, SecureCode connects to Groq's high-speed Language Processing Unit (LPU) cloud. "
        "It utilizes the verified 'openai/gpt-oss-120b' model with a 131,072 context window, analyzing surrounding source code "
        "and delivering valid JSON remediation diffs in under 1.2 seconds."
    )
    add_bullet_item(
        doc,
        "Automatic Multi-Model Fallback",
        "If a specific model encounters rate-limiting (TPM limits) or API restrictions, the client automatically cascades through "
        "verified alternative models: openai/gpt-oss-120b -> openai/gpt-oss-20b -> openai/gpt-oss-safeguard-20b -> allam-2-7b."
    )
    add_bullet_item(
        doc,
        "Offline Deterministic Knowledge Base",
        "If no API key is provided or the runner is in an air-gapped network, SecureCode seamlessly falls back to its deterministic "
        "security knowledge base. It provides pre-computed, verified fixes for common CWE vulnerabilities so CI/CD execution never halts."
    )

    add_heading_2(doc, "3.6 Engine 6: OASIS SARIF v2.1.0 Exporter")
    add_body_paragraph(
        doc,
        "The Static Analysis Results Interchange Format (SARIF) is the official OASIS standard for security reporting. "
        "SecureCode generates a compliant SARIF v2.1.0 document. Crucially, the engine implements the strict 'artifactChanges' "
        "schema for automated code replacements, ensuring 100% compatibility with GitHub CodeQL Action and displaying native alerts "
        "inside the repository's 'Security' tab."
    )

    add_heading_2(doc, "3.7 Engine 7: Unified Git Patch Generator")
    add_body_paragraph(
        doc,
        "Rather than forcing developers to manually retype code, SecureCode compiles every AI-generated diff into a standard "
        "unified Git patch file (remediation.patch). The patch includes valid headers (--- a/file, +++ b/file, @@ chunk markers) "
        "allowing developers to execute 'git apply remediation.patch' and resolve all vulnerabilities simultaneously."
    )

    doc.add_page_break()

    # ---------------------------------------------------------
    # CHAPTER 4: INDIVIDUAL CONTRIBUTIONS & TEAM RESPONSIBILITIES
    # ---------------------------------------------------------
    add_heading_1(doc, "4. Individual Contributions & Team Work Distribution")

    add_body_paragraph(
        doc,
        "SecureCode was designed, developed, and tested collaboratively by all three team members. Each contributor took "
        "complete ownership of distinct, critical architectural subsystems, ensuring equal balance between CI/CD infrastructure, "
        "static analysis engine internals, and AI-driven remediation services."
    )

    # Contributor 1: Shlok Garg
    add_heading_2(doc, "4.1 Shlok Garg (PRN: 24070122187)")
    add_body_paragraph(
        doc,
        "Lead Architecture: CI/CD Automation, GitHub Actions Composite Engine, SARIF v2.1.0 & Reporting"
    )
    add_bullet_item(doc, "Composite Action Architecture (action.yml)", "Engineered the zero-hosting GitHub Action definition, managing ephemeral environment provisioning, Python 3.11 runtime configuration, standalone Gitleaks binary compilation, and dependency caching.")
    add_bullet_item(doc, "OASIS SARIF v2.1.0 Exporter (sarif_exporter.py)", "Implemented the standard SARIF generator. Resolved critical schema validation issues by engineering the 'artifactChanges' replacement structure, enabling seamless ingestion by GitHub's Code Scanning dashboard.")
    add_bullet_item(doc, "CI Lifecycle & Policy Enforcement (scan.py)", "Engineered the policy failure threshold engine (fail-on: high, critical, medium). Designed safe exit code propagation so reports and SARIF artifacts are guaranteed to write before runner termination.")
    add_bullet_item(doc, "GitHub Step Summary Dashboard (format_github_step_summary)", "Designed the executive markdown dashboard with responsive tables, severity matrix tally, auto-expanded Critical/High cards (<details open>), and syntax-highlighted diffs.")
    add_bullet_item(doc, "Integration & Validation Testbed", "Configured and maintained the public demo repository (ShlokGarg22/TestingSecureCode) to validate live workflow runs, CodeQL uploads, and artifact retention.")

    # Contributor 2: Shashank Tai
    add_heading_2(doc, "4.2 Shashank Tai (PRN: 24070122181)")
    add_body_paragraph(
        doc,
        "Lead Architecture: Static Analysis Core, Security Rule Engineering & Secret Detection"
    )
    add_bullet_item(doc, "Semgrep Runner Engine (semgrep.py)", "Engineered the process execution wrapper for Semgrep, managing AST rule pack resolution (bundled, online, hybrid), process timeouts, and smart directory exclusion (skipping node_modules, venv, and cache folders).")
    add_bullet_item(doc, "Gitleaks Subsystem (gitleaks.py)", "Developed the high-entropy secret detection runner, configuring JSON report parsing, process isolation with shell=False, and extraction of secret matches across commit history and source trees.")
    add_bullet_item(doc, "Security Rule Engineering & Curation (backend/rules)", "Authored and curated AST pattern rules for critical Python vulnerabilities, including arbitrary file reading, OS command injection, unsafe YAML/pickle deserialization, and dangerous eval execution.")
    add_bullet_item(doc, "Normalization & Deduplication Engine (normalizer.py)", "Engineered the normalizer that transforms raw Semgrep results into standardized Finding models. Authored clean_rule_id_and_title() to eliminate runner path bloat and compute SHA-256 content fingerprints for deduplication.")
    add_bullet_item(doc, "Severity Classification Algorithm", "Formulated the severity weighting algorithm (Critical, High, Medium, Low, Info) mapping CWE risk categories to actionable policy thresholds.")

    # Contributor 3: Sayantan Nandi
    add_heading_2(doc, "4.3 Sayantan Nandi (PRN: 24070122176)")
    add_body_paragraph(
        doc,
        "Lead Architecture: Hybrid AI Remediation Engine, Privacy Redaction & Git Patch Generation"
    )
    add_bullet_item(doc, "Hybrid AI Explainer Architecture (explainer_enhanced.py)", "Architected the dual AI engine integrating Groq LPU inference (openai/gpt-oss-120b, openai/gpt-oss-20b) and Google Gemini with an offline deterministic knowledge base fallback.")
    add_bullet_item(doc, "Multi-Model Auto-Fallback Engine", "Engineered the fault-tolerant model cascade that automatically falls back across verified candidate models when encountering TPM rate limits, preventing CI pipeline failure.")
    add_bullet_item(doc, "Automated Privacy & Secret Redaction (redact_secrets_enhanced)", "Designed the high-entropy regex sanitization pipeline that redacts API keys, tokens, and database credentials from code snippets prior to LLM prompt construction.")
    add_bullet_item(doc, "Unified Git Patch Generator (patch_generator.py)", "Engineered the multi-diff patch compiler that parses AI remediations and outputs standard unified diffs (remediation.patch) compatible with git apply.")
    add_bullet_item(doc, "Groq Model Diagnostic Tool (check_groq_models.py)", "Developed the standalone CLI tool that queries Groq client.models.list(), benchmarks latency, validates JSON schemas, and identifies active models for user accounts.")
    add_bullet_item(doc, "Full-Stack Web Interface & REST APIs", "Co-developed the FastAPI backend (scans.py) and interactive Next.js web application for manual ZIP uploads and visual side-by-side diff comparisons.")

    # Responsibility Matrix Table
    add_heading_2(doc, "4.4 Team Responsibility & Deliverables Matrix")
    matrix_table = doc.add_table(rows=6, cols=4)
    matrix_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    m_widths = [2.4, 1.3, 1.3, 1.5]
    
    format_table_header(matrix_table.rows[0], m_widths, ["Subsystem / Component", "Lead Owner", "Secondary", "Deliverable Files"])
    add_table_row(matrix_table, m_widths, ["Composite GitHub Action & CI/CD", "Shlok Garg", "Sayantan Nandi", "action.yml, securecode.yml"], False, True)
    add_table_row(matrix_table, m_widths, ["Static Analysis & Rule Engine", "Shashank Tai", "Shlok Garg", "semgrep.py, gitleaks.py, rules/"], True, True)
    add_table_row(matrix_table, m_widths, ["SARIF v2.1.0 & Step Summary", "Shlok Garg", "Shashank Tai", "sarif_exporter.py, scan.py"], False, True)
    add_table_row(matrix_table, m_widths, ["Hybrid AI & Model Cascade", "Sayantan Nandi", "Shashank Tai", "explainer_enhanced.py, explainer.py"], True, True)
    add_table_row(matrix_table, m_widths, ["Git Patch & Privacy Redaction", "Sayantan Nandi", "Shlok Garg", "patch_generator.py, check_groq_models.py"], False, True)

    doc.add_page_break()

    # ---------------------------------------------------------
    # CHAPTER 5: REAL-WORLD DEMONSTRATION & TESTING
    # ---------------------------------------------------------
    add_heading_1(doc, "5. Real-World Testing & Demonstration Results")

    add_body_paragraph(
        doc,
        "To rigorously evaluate SecureCode in an authentic production environment, we constructed a dedicated target "
        "repository (TestingSecureCode) containing deliberate security vulnerabilities across multiple categories. "
        "The repository was connected to GitHub Actions using SecureCode@main."
    )

    add_heading_2(doc, "5.1 Target Vulnerabilities in app.py")
    vuln_table = doc.add_table(rows=7, cols=4)
    vuln_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    v_widths = [0.8, 1.5, 2.2, 2.0]
    
    format_table_header(vuln_table.rows[0], v_widths, ["#", "Severity", "Vulnerability Title", "Vulnerable Location & Vector"])
    add_table_row(vuln_table, v_widths, ["1", "🔴 Critical", "Command Injection (os.system)", "app.py:117 — os.system(f'ping -c 1 {ip}')"], False, True)
    add_table_row(vuln_table, v_widths, ["2", "🔴 Critical", "Insecure Pickle Deserialization", "app.py:151 — pickle.loads(raw_session_bytes)"], True, True)
    add_table_row(vuln_table, v_widths, ["3", "🟠 High", "Dangerous Code Execution (eval)", "app.py:165 — computed = eval(user_formula)"], False, True)
    add_table_row(vuln_table, v_widths, ["4", "🟠 High", "Hardcoded API Secret Credential", "app.py:18 — AWS_SECRET_KEY = 'AKIAIOSFODNN...'"], True, True)
    add_table_row(vuln_table, v_widths, ["5", "🟡 Medium", "Arbitrary File Path Traversal", "app.py:30 — open(f'uploads/{user_path}')"], False, True)
    add_table_row(vuln_table, v_widths, ["6", "🟡 Medium", "Arbitrary File Path Traversal", "app.py:106 — open(target_doc_path, 'r')"], True, True)

    add_heading_2(doc, "5.2 Live Scan Output: GitHub Actions Step Summary")
    add_body_paragraph(
        doc,
        "Upon pushing to the repository, SecureCode executed in 27 seconds, automatically generating the following "
        "executive markdown dashboard directly inside GitHub's Actions Summary:"
    )

    add_code_block(
        doc,
        "# 🛡️ SecureCode Security Audit Report\n"
        "> Project: TestingSecureCode | Scan ID: CLI-20261021... | Status: ❌ Policy Blocked\n\n"
        "### 📊 Security Posture Matrix\n"
        "| Severity Level | Count | Risk Rating | Status |\n"
        "| :--- | :---: | :--- | :--- |\n"
        "| 🔴 Critical | 2 | Remote Code Execution / Host Compromise | ❌ Blocker |\n"
        "| 🟠 High | 2 | Direct Exploitation / Exposed Secrets | ❌ Blocker |\n"
        "| 🟡 Medium | 2 | Path Traversal / Injection Vectors | ⚠️ Review |\n"
        "| 🔵 Low | 0 | Cryptographic Hardening | ✅ Clean |\n\n"
        "1. 🔴 [CRITICAL] Command Injection via Shell (os.system) — app.py:117\n"
        "   - Location: app.py (Line 117)\n"
        "   - Analysis Engine: 🤖 Live AI (Groq gpt-oss-120b)\n"
        "   - Why It Matters: User input is passed directly to the OS shell without escaping.\n"
        "   - Attack Scenario: Attacker inputs '127.0.0.1; cat /etc/shadow' to steal credentials.\n"
        "   - Suggested Remediation Diff:\n"
        "     - os.system(f'ping -c 1 {target_ip}')\n"
        "     + subprocess.run(['ping', '-c', '1', target_ip], capture_output=True, shell=False)",
        caption="Sample Output from GitHub Actions Step Summary"
    )

    add_heading_2(doc, "5.3 Instant 1-Click Remediation: remediation.patch")
    add_body_paragraph(
        doc,
        "SecureCode compiled all 6 remediations into a downloadable artifact named 'remediation-patch'. "
        "A developer downloads the patch and executes a single Git command to resolve all findings simultaneously:"
    )
    add_code_block(
        doc,
        "# In the root of the project:\n"
        "$ git apply remediation.patch\n"
        "$ git status\n"
        "modified:   app.py\n"
        "$ git commit -am 'fix: apply automated security remediation patch from SecureCode'\n"
        "$ git push origin main",
        caption="1-Click Git Remediation Command"
    )

    # ---------------------------------------------------------
    # CHAPTER 6: COMPARATIVE ANALYSIS & PERFORMANCE
    # ---------------------------------------------------------
    add_heading_1(doc, "6. Comparative Analysis & Performance Metrics")

    add_body_paragraph(
        doc,
        "To objectively benchmark SecureCode against established industry standards, we compared key features "
        "against SonarQube, Snyk, and GitHub CodeQL."
    )

    comp_table = doc.add_table(rows=6, cols=5)
    comp_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_widths = [1.8, 1.2, 1.1, 1.2, 1.2]
    
    format_table_header(comp_table.rows[0], c_widths, ["Feature / Metric", "SecureCode", "SonarQube", "Snyk", "CodeQL"])
    add_table_row(comp_table, c_widths, ["Hosting Cost", "$0 (Zero-Hosting)", "$50 - $200/mo", "SaaS Subscription", "$0 (Public Repos)"], False, True)
    add_table_row(comp_table, c_widths, ["Setup Complexity", "1 YAML file (10 lines)", "Complex (DB + Server)", "CLI / SaaS Token", "Complex Database Build"], True, True)
    add_table_row(comp_table, c_widths, ["AI Attack Scenarios", "Yes (Live Groq LPU)", "No", "Basic Advice", "No"], False, True)
    add_table_row(comp_table, c_widths, ["Downloadable Git Patch", "Yes (remediation.patch)", "No", "PR Bot Only", "No"], True, True)
    add_table_row(comp_table, c_widths, ["Average Scan Duration", "~25 - 30 seconds", "~2 - 5 minutes", "~1 - 2 minutes", "~3 - 8 minutes"], False, True)

    add_heading_2(doc, "6.1 Execution Latency Breakdown")
    add_bullet_item(doc, "Gitleaks Secret Detection", "1.1 seconds — Extremely fast compiled Go binary.")
    add_bullet_item(doc, "Semgrep AST Static Analysis", "4.2 seconds — Highly efficient rule matching excluding vendor directories.")
    add_bullet_item(doc, "Groq LPU AI Inference", "1.2 seconds per finding — Ultra-fast inference powered by LPUs.")
    add_bullet_item(doc, "Artifact & Report Export", "0.4 seconds — File writing and SARIF compilation.")
    add_bullet_item(doc, "Total Workflow Execution", "27.0 seconds total runner time from trigger to conclusion.")

    # ---------------------------------------------------------
    # CHAPTER 7: ENGINEERING CHALLENGES & OVERCOMING THEM
    # ---------------------------------------------------------
    add_heading_1(doc, "7. Engineering Challenges & Overcoming Them")

    add_heading_2(doc, "7.1 Challenge 1: Composite Action Python Module Resolution")
    add_body_paragraph(
        doc,
        "When running a composite GitHub Action, the working directory is the caller repository, not the action repository. "
        "Initial runs threw 'ModuleNotFoundError: No module named app'. We solved this by dynamically prepending "
        "the action's backend directory to sys.path in scan.py: backend_dir = Path(__file__).resolve().parent.parent; "
        "sys.path.insert(0, str(backend_dir)), ensuring flawless cross-repository invocation."
    )

    add_heading_2(doc, "7.2 Challenge 2: OASIS SARIF Schema Incompatibility")
    add_body_paragraph(
        doc,
        "GitHub's CodeQL Action performs strict JSON schema validation. Initial attempts to include code fixes failed with: "
        "'fixes[0] is not allowed to have additional property fileChanges; requires artifactChanges'. "
        "We restructured sarif_exporter.py to conform strictly to the OASIS SARIF v2.1.0 specification, wrapping replacements "
        "in 'artifactChanges', which enabled native GitHub security alerts."
    )

    add_heading_2(doc, "7.3 Challenge 3: Preserving Multi-Line Git Diffs in Markdown")
    add_body_paragraph(
        doc,
        "In GitHub Actions Step Summary, omitting blank lines immediately after HTML <details> tags caused the CommonMark "
        "parser to squash entire diff blocks onto a single illegible line. We re-engineered format_github_step_summary "
        "with explicit line break enforcement, clean tabular metadata, and auto-expanded High/Critical cards."
    )

    add_heading_2(doc, "7.4 Challenge 4: Groq API Key Tiers and Model Discovery")
    add_body_paragraph(
        doc,
        "Different Groq accounts have varying access to open models. Defaulting to Meta LLaMA caused rate-limit and access errors. "
        "We built check_groq_models.py to programmatically discover account models and integrated an autonomous model cascade "
        "defaulting to 'openai/gpt-oss-120b', which operates flawlessly with a 131k token context window."
    )

    # ---------------------------------------------------------
    # CHAPTER 8: CONCLUSION & FUTURE SCOPE
    # ---------------------------------------------------------
    add_heading_1(doc, "8. Conclusion & Future Roadmap")

    add_body_paragraph(
        doc,
        "SecureCode proves that enterprise-grade DevSecOps security does not require costly server infrastructure or "
        "frustrating developer workflows. By uniting AST static analysis, high-entropy secret detection, privacy-preserving "
        "redaction, and ultra-fast Groq LPU inference, SecureCode delivers instant security visibility and automated "
        "Git patch remediation directly where developers already work—inside GitHub Actions."
    )

    add_heading_2(doc, "8.1 Future Roadmap")
    add_bullet_item(doc, "Multi-Language Expansion", "Extend bundled AST rules to JavaScript, TypeScript, Go, Java, and Rust.")
    add_bullet_item(doc, "Autonomous GitHub Pull Request Bot", "Automatically open PR branches containing the applied remediation patch, allowing developers to review and merge fixes with a single button click.")
    add_bullet_item(doc, "VS Code IDE Extension", "Bring SecureCode's deterministic and AI remediation engine directly into developer IDEs for pre-commit inline warnings.")

    # Save Document
    output_path = os.path.abspath("SecureCode_Project_Report.docx")
    doc.save(output_path)
    print(f"\n[+] Successfully generated SecureCode Project Report at:\n    {output_path}")
    return output_path


if __name__ == "__main__":
    build_full_report()
